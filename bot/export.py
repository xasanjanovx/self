"""Экспорт операций в Excel (openpyxl). Выполнять в потоке — это CPU/IO."""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import categories as cats
from . import finance as fin

_HEAD_FILL = PatternFill("solid", fgColor="DDE7F5")


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(48, max(10, width + 2))


def build_xlsx(entries: list[dict[str, Any]], *, period: fin.Period, lang: str, currency: str) -> bytes:
    rows = fin.entries_between(entries, period.start, period.end)
    rows = sorted(rows, key=lambda r: (str(r.get("entry_date")), str(r.get("created_at"))))
    stats = fin.compute_stats(entries, period)
    uz = lang == "uz"

    wb = Workbook()
    ws = wb.active
    ws.title = "Operatsiyalar" if uz else "Операции"
    head = ["Sana", "Tur", "Kategoriya", "Summa", "Hisob", "Izoh"] if uz else ["Дата", "Тип", "Категория", "Сумма", "Счёт", "Заметка"]
    ws.append(head)
    for r in rows:
        transfer = fin.transfer_from_note(r.get("note"))
        if transfer:
            kind = "O'tkazma" if uz else "Перевод"
            account = fin.transfer_label(transfer[0], transfer[1], lang)
            amount = float(r.get("amount") or 0)
        else:
            income = r.get("entry_type") == "income"
            kind = ("Kirim" if uz else "Доход") if income else ("Chiqim" if uz else "Расход")
            account = fin.bucket_label(fin.bucket_from_note(r.get("note")), lang)
            amount = float(r.get("amount") or 0) * (1 if income else -1)
        ws.append([
            date.fromisoformat(str(r.get("entry_date"))[:10]),
            kind,
            cats.label(fin.entry_category_key(r), lang, with_emoji=False),
            amount,
            account,
            fin.clean_note(r.get("note")) or "",
        ])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = _HEAD_FILL
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[3].number_format = "#,##0"
    ws.freeze_panes = "A2"
    _autosize(ws)

    ws2 = wb.create_sheet("Toifalar" if uz else "По категориям")
    ws2.append(["Kategoriya", "Summa", "Ulush", "Soni"] if uz else ["Категория", "Сумма", "Доля", "Операций"])
    total = stats.expense or 1.0
    for key, amount, count in stats.by_category:
        ws2.append([cats.label(key, lang, with_emoji=False), amount, amount / total, count])
    ws2.append([])
    ws2.append(["Kirimlar" if uz else "Доходы"])
    for key, amount, count in stats.income_by_category:
        ws2.append([cats.label(key, lang, with_emoji=False), amount, "", count])
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = _HEAD_FILL
    for row in ws2.iter_rows(min_row=2):
        row[1].number_format = "#,##0"
        row[2].number_format = "0%"
    _autosize(ws2)

    ws3 = wb.create_sheet("Natija" if uz else "Итоги")
    ws3.append(["Davr" if uz else "Период", f"{period.start:%d.%m.%Y} – {period.end:%d.%m.%Y}"])
    ws3.append(["Valyuta" if uz else "Валюта", currency])
    ws3.append(["Chiqim" if uz else "Расход", stats.expense])
    ws3.append(["Kirim" if uz else "Доход", stats.income])
    ws3.append(["Natija" if uz else "Итог", stats.net])
    ws3.append(["Kuniga o'rtacha" if uz else "Средний расход в день", stats.avg_per_day])
    ws3.append(["Operatsiyalar" if uz else "Операций", stats.ops])
    for row in ws3.iter_rows(min_row=3, max_row=6):
        row[1].number_format = "#,##0"
    for row in ws3.iter_rows():
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(horizontal="left")
    _autosize(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
